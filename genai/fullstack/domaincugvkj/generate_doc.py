"""
Fullstack Exam Document Generator
Generates a comprehensive Word document with Aim, Algorithm, Workflow, Procedure, Code, Files, and How to Run
for all 4 questions (SET A Q1, SET A Q2, SET B Q1, SET B Q2).
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

def set_cell_shading(cell, color_hex):
    """Set background color for a table cell."""
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def add_code_block(doc, code_text, filename=""):
    """Add a formatted code block to the document."""
    if filename:
        p = doc.add_paragraph()
        run = p.add_run(f"File: {filename}")
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0, 100, 0)

    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.cell(0, 0)
    set_cell_shading(cell, 'F5F5F5')
    cell.text = ""
    para = cell.paragraphs[0]
    run = para.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 30, 30)
    doc.add_paragraph("")  # spacing

def add_heading_styled(doc, text, level=1):
    """Add a styled heading."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0, 51, 102)

def add_bullet_list(doc, items):
    """Add a bulleted list."""
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

def add_numbered_list(doc, items):
    """Add a numbered list."""
    for item in items:
        doc.add_paragraph(item, style='List Number')

def add_page_break(doc):
    doc.add_page_break()

def add_image_with_caption(doc, image_path, caption):
    """Add an image with a center-aligned caption to the document."""
    if os.path.exists(image_path):
        doc.add_paragraph()
        run = doc.add_paragraph().add_run()
        run.add_picture(image_path, width=Inches(6.0))
        cap_p = doc.add_paragraph(caption)
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cap_p.runs:
            r.font.italic = True

# ===============================================================
# MAIN DOCUMENT GENERATION
# ===============================================================

doc = Document()

# --- Title Page ---
doc.add_paragraph("")
doc.add_paragraph("")
doc.add_paragraph("")
title = doc.add_heading("FULLSTACK DEVELOPMENT", level=0)
for run in title.runs:
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0, 51, 102)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_heading("Lab Examination", level=1)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in subtitle.runs:
    run.font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph("")
desc = doc.add_paragraph("Comprehensive documentation covering Angular and Node.js/Express applications")
desc.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("")
doc.add_paragraph("")

# Contents table
contents_table = doc.add_table(rows=5, cols=2)
contents_table.style = 'Table Grid'
contents_table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = contents_table.rows[0].cells
headers[0].text = "Set / Question"
headers[1].text = "Topic"
set_cell_shading(headers[0], '003366')
set_cell_shading(headers[1], '003366')
for cell in headers:
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True

data = [
    ("SET A - Question 1", "Clock Angle Calculator (Angular)"),
    ("SET A - Question 2", "CRM Restaurant Application (Node/Express)"),
    ("SET B - Question 1", "Kids Library Gift Coupon System (Angular)"),
    ("SET B - Question 2", "CRM Restaurant Application (Node/Express)"),
]
for i, (col1, col2) in enumerate(data):
    row = contents_table.rows[i + 1].cells
    row[0].text = col1
    row[1].text = col2

add_page_break(doc)

# ================================================================
# SET A - QUESTION 1: Clock Angle Calculator
# ================================================================

add_heading_styled(doc, "SET A - QUESTION 1", level=1)
add_heading_styled(doc, "Clock Angle Calculator using Angular", level=2)

# --- AIM ---
add_heading_styled(doc, "1. Aim", level=2)
doc.add_paragraph(
    "To design and develop a web application using Angular framework that computes the angle "
    "between the hour and minute hands of a clock. The application performs appropriate form handling "
    "by fetching hour and minute values from respective textboxes and displays the calculated angle "
    "in a third textbox upon clicking a button. The angle is computed using the standard clock angle formula."
)

# --- ALGORITHM ---
add_heading_styled(doc, "2. Algorithm", level=2)
add_numbered_list(doc, [
    "START",
    "Accept the hour value (H) from the user (1 to 12).",
    "Accept the minute value (M) from the user (0 to 59).",
    "Validate the input: H must be between 1-12, M must be between 0-59.",
    "Calculate the hour hand angle: HourAngle = (H % 12) × 30 + M × 0.5",
    "Calculate the minute hand angle: MinuteAngle = M × 6",
    "Calculate the difference: Diff = |HourAngle - MinuteAngle|",
    "If Diff > 180, then Angle = 360 - Diff, else Angle = Diff",
    "Display the computed angle in the result textbox.",
    "STOP"
])

doc.add_paragraph("")
doc.add_paragraph("Formula Explanation:")
doc.add_paragraph("• The hour hand moves 360°/12 = 30° per hour and 0.5° per minute.", style='List Bullet')
doc.add_paragraph("• The minute hand moves 360°/60 = 6° per minute.", style='List Bullet')
doc.add_paragraph("• The angle between the two hands is the absolute difference, taking the minimum of the two possible arcs.", style='List Bullet')

# --- WORKFLOW ---
add_heading_styled(doc, "3. Workflow", level=2)
add_numbered_list(doc, [
    "User opens the Angular application in the browser.",
    "The form displays two input textboxes for Hour and Minute.",
    "User enters the hour value (1-12) in the first textbox.",
    "User enters the minute value (0-59) in the second textbox.",
    "User clicks the 'Calculate Angle' button.",
    "Angular performs form validation (checks if values are within valid ranges).",
    "If validation fails, an error message is displayed.",
    "If validation passes, the component calculates the angle using the formula.",
    "The computed angle is displayed in a read-only third textbox.",
    "User can click 'Reset' to clear all fields and start over."
])

# --- PROCEDURE ---
add_heading_styled(doc, "4. Procedure", level=2)
add_numbered_list(doc, [
    "Install Node.js and npm on the system.",
    "Install Angular CLI globally using: npm install -g @angular/cli",
    "Create a new Angular project: ng new clock-angle-calculator",
    "Navigate to the project directory: cd clock-angle-calculator",
    "Import FormsModule in app.module.ts for two-way data binding (ngModel).",
    "Create the AppComponent with properties: hour, minute, angle, errorMessage.",
    "Implement the calculateAngle() method with the clock angle formula.",
    "Implement the reset() method to clear all form fields.",
    "Design the HTML template with input fields, buttons, and result display.",
    "Add CSS styling for a clean and user-friendly interface.",
    "Run the application using: ng serve",
    "Open browser and navigate to http://localhost:4200",
    "Test with various hour and minute values to verify correctness."
])

# --- CODE ---
add_heading_styled(doc, "5. Code", level=2)

# app.module.ts
add_code_block(doc, """import { NgModule } from '@angular/core';
import { BrowserModule } from '@angular/platform-browser';
import { FormsModule } from '@angular/forms';
import { AppComponent } from './app.component';

@NgModule({
  declarations: [AppComponent],
  imports: [BrowserModule, FormsModule],
  providers: [],
  bootstrap: [AppComponent]
})
export class AppModule {}""", "src/app/app.module.ts")

# app.component.ts
add_code_block(doc, """import { Component } from '@angular/core';

@Component({
  selector: 'app-root',
  templateUrl: './app.component.html',
  styleUrls: ['./app.component.css']
})
export class AppComponent {
  title = 'Clock Angle Calculator';
  hour: number | null = null;
  minute: number | null = null;
  angle: number | null = null;
  errorMessage: string = '';

  calculateAngle(): void {
    this.errorMessage = '';
    this.angle = null;

    // Form validation
    if (this.hour === null || this.minute === null) {
      this.errorMessage = 'Please enter both hour and minute values.';
      return;
    }
    if (this.hour < 1 || this.hour > 12) {
      this.errorMessage = 'Hour must be between 1 and 12.';
      return;
    }
    if (this.minute < 0 || this.minute > 59) {
      this.errorMessage = 'Minute must be between 0 and 59.';
      return;
    }

    // Calculate the angle of the hour hand from 12 o'clock position
    // Hour hand moves 0.5 degrees per minute (360/12/60)
    const hourAngle = (this.hour % 12) * 30 + this.minute * 0.5;

    // Calculate the angle of the minute hand from 12 o'clock position
    // Minute hand moves 6 degrees per minute (360/60)
    const minuteAngle = this.minute * 6;

    // Find the absolute difference
    let diff = Math.abs(hourAngle - minuteAngle);

    // Take the smaller angle (clock has two arcs)
    if (diff > 180) {
      diff = 360 - diff;
    }

    this.angle = diff;
  }

  reset(): void {
    this.hour = null;
    this.minute = null;
    this.angle = null;
    this.errorMessage = '';
  }
}""", "src/app/app.component.ts")

# app.component.html
add_code_block(doc, """<div class="container">
  <h1>{{ title }}</h1>
  <div class="form-group">
    <label for="hour">Hour (1-12):</label>
    <input type="number" id="hour" [(ngModel)]="hour"
           placeholder="Enter hour" min="1" max="12" />
  </div>
  <div class="form-group">
    <label for="minute">Minute (0-59):</label>
    <input type="number" id="minute" [(ngModel)]="minute"
           placeholder="Enter minute" min="0" max="59" />
  </div>
  <div class="button-group">
    <button (click)="calculateAngle()">Calculate Angle</button>
    <button (click)="reset()" class="reset-btn">Reset</button>
  </div>
  <div class="form-group" *ngIf="angle !== null">
    <label for="result">Angle between hands:</label>
    <input type="text" id="result" [value]="angle + ' degrees'" readonly />
  </div>
  <div class="error" *ngIf="errorMessage">{{ errorMessage }}</div>
</div>""", "src/app/app.component.html")

# app.component.css
add_code_block(doc, """.container {
  max-width: 400px;
  margin: 50px auto;
  padding: 30px;
  border: 1px solid #ccc;
  border-radius: 10px;
  font-family: Arial, sans-serif;
  box-shadow: 0 2px 10px rgba(0, 0, 0, 0.1);
}

h1 {
  text-align: center;
  color: #333;
  margin-bottom: 20px;
}

.form-group { margin-bottom: 15px; }

label {
  display: block;
  margin-bottom: 5px;
  font-weight: bold;
  color: #555;
}

input {
  width: 100%;
  padding: 10px;
  border: 1px solid #ccc;
  border-radius: 5px;
  font-size: 16px;
  box-sizing: border-box;
}

input[readonly] {
  background-color: #e9ecef;
  font-weight: bold;
  color: #28a745;
}

.button-group {
  display: flex;
  gap: 10px;
  margin-bottom: 15px;
}

button {
  flex: 1;
  padding: 10px;
  background-color: #007bff;
  color: white;
  border: none;
  border-radius: 5px;
  font-size: 16px;
  cursor: pointer;
}

button:hover { background-color: #0056b3; }

.reset-btn { background-color: #6c757d; }
.reset-btn:hover { background-color: #545b62; }

.error {
  color: red;
  font-size: 14px;
  margin-top: 10px;
  text-align: center;
}""", "src/app/app.component.css")

# main.ts
add_code_block(doc, """import { platformBrowserDynamic } from '@angular/platform-browser-dynamic';
import { AppModule } from './app/app.module';

platformBrowserDynamic().bootstrapModule(AppModule)
  .catch(err => console.error(err));""", "src/main.ts")

# index.html
add_code_block(doc, """<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <title>Clock Angle Calculator</title>
  <base href="/">
  <meta name="viewport" content="width=device-width, initial-scale=1">
</head>
<body>
  <app-root></app-root>
</body>
</html>""", "src/index.html")

# --- FILES ---
add_heading_styled(doc, "6. Files", level=2)
files_table = doc.add_table(rows=7, cols=3)
files_table.style = 'Table Grid'
file_headers = files_table.rows[0].cells
file_headers[0].text = "S.No"
file_headers[1].text = "File Name"
file_headers[2].text = "Description"
set_cell_shading(file_headers[0], '003366')
set_cell_shading(file_headers[1], '003366')
set_cell_shading(file_headers[2], '003366')
for cell in file_headers:
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True

file_data = [
    ("1", "src/app/app.module.ts", "Root module that imports FormsModule for ngModel binding"),
    ("2", "src/app/app.component.ts", "Component with clock angle calculation logic and form validation"),
    ("3", "src/app/app.component.html", "HTML template with input fields and result display"),
    ("4", "src/app/app.component.css", "CSS styling for the form and layout"),
    ("5", "src/main.ts", "Bootstrap file that initializes the Angular application"),
    ("6", "src/index.html", "Root HTML file with app-root selector"),
]
for i, (sno, fname, desc) in enumerate(file_data):
    row = files_table.rows[i + 1].cells
    row[0].text = sno
    row[1].text = fname
    row[2].text = desc

# --- HOW TO RUN ---
add_heading_styled(doc, "7. How to Run", level=2)
add_numbered_list(doc, [
    "Ensure Node.js (v14+) and npm are installed on the system.",
    "Install Angular CLI globally:\n   npm install -g @angular/cli",
    "Create a new Angular project:\n   ng new clock-angle-calculator",
    "Navigate to the project folder:\n   cd clock-angle-calculator",
    "Replace the files in src/app/ with the code files provided above.",
    "Run the development server:\n   ng serve",
    "Open a web browser and navigate to:\n   http://localhost:4200",
    "Enter hour (1-12) and minute (0-59) values and click 'Calculate Angle'.",
    "The angle between the hour and minute hands will be displayed in the result textbox."
])

doc.add_paragraph("")
p = doc.add_paragraph()
run = p.add_run("Sample Test Cases:")
run.bold = True
run.font.size = Pt(11)

test_table = doc.add_table(rows=6, cols=3)
test_table.style = 'Table Grid'
test_headers = test_table.rows[0].cells
test_headers[0].text = "Hour"
test_headers[1].text = "Minute"
test_headers[2].text = "Expected Angle"
set_cell_shading(test_headers[0], '4CAF50')
set_cell_shading(test_headers[1], '4CAF50')
set_cell_shading(test_headers[2], '4CAF50')
for cell in test_headers:
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True

test_data = [
    ("3", "0", "90°"),
    ("6", "0", "180°"),
    ("12", "30", "165°"),
    ("9", "0", "90°"),  # fixed: 270 -> 360-270=90
    ("3", "30", "75°"),
]
for i, (h, m, a) in enumerate(test_data):
    row = test_table.rows[i + 1].cells
    row[0].text = h
    row[1].text = m
    row[2].text = a

add_heading_styled(doc, "8. Output", level=2)
add_image_with_caption(doc, r"c:\RISK_ASSESSMENT\genai\fullstack\screenshots\setA_q1_output.png", "Clock Angle Calculator Output")

add_page_break(doc)

# ================================================================
# SET A - QUESTION 2: CRM Restaurant (Node/Express)
# ================================================================

add_heading_styled(doc, "SET A - QUESTION 2", level=1)
add_heading_styled(doc, "CRM Application for Restaurant using Node.js and Express", level=2)

# --- AIM ---
add_heading_styled(doc, "1. Aim", level=2)
doc.add_paragraph(
    "To design and develop a Customer Relationship Management (CRM) application for a restaurant "
    "using Node.js and Express framework. The application implements user authentication with session "
    "management, a home page displaying welcome message and menu card image, a surprise offer page, "
    "customer information collection and display in a table, and logout functionality that destroys "
    "the session across all pages."
)

# --- ALGORITHM ---
add_heading_styled(doc, "2. Algorithm", level=2)
add_numbered_list(doc, [
    "START",
    "Initialize Express application with session middleware and EJS view engine.",
    "Define hardcoded credentials (username: admin, password: password123).",
    "Route GET /login: Render the login page with username and password textboxes.",
    "Route POST /home: Receive form data, validate credentials against hardcoded values.",
    "If credentials are valid, create a session and render home page with 'Welcome <<username>>'.",
    "If credentials are invalid, re-render login page with error message.",
    "Home page displays: menu card image (as hyperlink to /surprise), customer form (age & mobile textboxes with 'add' hyperlink), and logout link.",
    "Route GET /surprise: Check session, display 'surprise offer of Rs.250 in the next order' in <h3> tag.",
    "Route GET /customer: Check session, receive age and mobile via query parameters, add to customer array, display all customers in a table.",
    "Route GET /logout: Destroy the session and redirect to login page.",
    "All pages (home, surprise, customer) include a logout hyperlink.",
    "STOP"
])

# --- WORKFLOW ---
add_heading_styled(doc, "3. Workflow", level=2)
add_numbered_list(doc, [
    "User accesses the application at http://localhost:3000.",
    "User is redirected to the login page (/login).",
    "User enters username and password and submits the form (POST to /home).",
    "Server validates credentials and creates a session on success.",
    "Home page displays 'Welcome admin' with menu card image and customer form.",
    "Clicking the menu card image navigates to /surprise page.",
    "Surprise page shows 'surprise offer of Rs.250 in the next order'.",
    "On home page, user enters customer age and mobile, clicks 'add' hyperlink.",
    "Customer route (/customer) displays the customer data in a table.",
    "Clicking 'Logout' on any page destroys the session and redirects to login."
])

# --- PROCEDURE ---
add_heading_styled(doc, "4. Procedure", level=2)
add_numbered_list(doc, [
    "Create a new project folder: mkdir crm-restaurant && cd crm-restaurant",
    "Initialize Node.js project: npm init -y",
    "Install dependencies: npm install express express-session ejs",
    "Create the main server file (app.js) with Express configuration.",
    "Configure session middleware with a secret key.",
    "Set EJS as the view engine.",
    "Create the views folder with EJS templates: login.ejs, home.ejs, surprise.ejs, customer.ejs.",
    "Create the public/images folder and add a menu card image.",
    "Define routes: GET /login, POST /home, GET /surprise, GET /customer, GET /logout.",
    "Add session checking middleware to protected routes.",
    "Run the server: node app.js",
    "Test the application at http://localhost:3000"
])

# --- CODE ---
add_heading_styled(doc, "5. Code", level=2)

# package.json
add_code_block(doc, """{
  "name": "crm-restaurant",
  "version": "1.0.0",
  "description": "CRM Application for Restaurant using Node and Express",
  "main": "app.js",
  "scripts": {
    "start": "node app.js"
  },
  "dependencies": {
    "ejs": "^3.1.9",
    "express": "^4.18.2",
    "express-session": "^1.17.3"
  }
}""", "package.json")

# app.js
add_code_block(doc, """const express = require('express');
const session = require('express-session');
const path = require('path');
const app = express();

// Middleware
app.use(express.urlencoded({ extended: true }));
app.use(express.static(path.join(__dirname, 'public')));

// Session configuration
app.use(session({
    secret: 'restaurant-crm-secret-key',
    resave: false,
    saveUninitialized: false,
    cookie: { maxAge: 600000 } // 10 minutes
}));

// Set EJS as view engine
app.set('view engine', 'ejs');
app.set('views', path.join(__dirname, 'views'));

// Hardcoded credentials
const VALID_USERNAME = 'admin';
const VALID_PASSWORD = 'password123';

// Store customer data in memory
let customers = [];

// GET /login - Display login page
app.get('/login', (req, res) => {
    res.render('login', { error: null });
});

// GET / - Redirect to login
app.get('/', (req, res) => {
    res.redirect('/login');
});

// POST /home - Handle login form submission
app.post('/home', (req, res) => {
    const { username, password } = req.body;
    if (username === VALID_USERNAME && password === VALID_PASSWORD) {
        req.session.username = username;
        req.session.isLoggedIn = true;
        res.render('home', { username: username });
    } else {
        res.render('login', { error: 'Invalid username or password!' });
    }
});

// GET /surprise - Display surprise offer
app.get('/surprise', (req, res) => {
    if (!req.session.isLoggedIn) {
        return res.redirect('/login');
    }
    res.render('surprise', { username: req.session.username });
});

// GET /customer - Display customer information in table
app.get('/customer', (req, res) => {
    if (!req.session.isLoggedIn) {
        return res.redirect('/login');
    }
    const { age, mobile } = req.query;
    if (age && mobile) {
        customers.push({ age: age, mobile: mobile });
    }
    res.render('customer', {
        customers: customers,
        username: req.session.username
    });
});

// GET /logout - Destroy session and redirect to login
app.get('/logout', (req, res) => {
    req.session.destroy((err) => {
        if (err) {
            console.log('Error destroying session:', err);
        }
        customers = [];
        res.redirect('/login');
    });
});

// Start server
const PORT = 3000;
app.listen(PORT, () => {
    console.log('Restaurant CRM Server running at http://localhost:' + PORT);
});""", "app.js")

# login.ejs
add_code_block(doc, """<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Restaurant CRM - Login</title>
    <style>
        body {
            font-family: Arial, sans-serif;
            background-color: #f4f4f4;
            display: flex;
            justify-content: center;
            align-items: center;
            height: 100vh;
            margin: 0;
        }
        .login-container {
            background: white;
            padding: 40px;
            border-radius: 10px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.1);
            width: 350px;
        }
        h2 { text-align: center; color: #333; }
        .form-group { margin-bottom: 15px; }
        label {
            display: block;
            margin-bottom: 5px;
            font-weight: bold;
        }
        input[type="text"], input[type="password"] {
            width: 100%;
            padding: 10px;
            border: 1px solid #ccc;
            border-radius: 5px;
            box-sizing: border-box;
        }
        button {
            width: 100%;
            padding: 12px;
            background-color: #FF5722;
            color: white;
            border: none;
            border-radius: 5px;
            font-size: 16px;
            cursor: pointer;
        }
        button:hover { background-color: #E64A19; }
        .error { color: red; text-align: center; }
    </style>
</head>
<body>
    <div class="login-container">
        <h2>Restaurant CRM - Login</h2>
        <%% if (error) { %>
            <p class="error"><%%= error %></p>
        <%% } %>
        <form action="/home" method="POST">
            <div class="form-group">
                <label for="username">Username:</label>
                <input type="text" id="username" name="username" required>
            </div>
            <div class="form-group">
                <label for="password">Password:</label>
                <input type="password" id="password" name="password" required>
            </div>
            <button type="submit">Login</button>
        </form>
    </div>
</body>
</html>""", "views/login.ejs")

# home.ejs
add_code_block(doc, """<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Restaurant CRM - Home</title>
    <style>
        body {
            font-family: Arial, sans-serif;
            background-color: #f4f4f4;
            padding: 20px;
        }
        .container {
            max-width: 700px;
            margin: 0 auto;
            background: white;
            padding: 30px;
            border-radius: 10px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.1);
        }
        h1 { color: #FF5722; text-align: center; }
        .menu-section { text-align: center; margin: 20px 0; }
        .menu-section img {
            width: 300px;
            border-radius: 10px;
            border: 2px solid #ddd;
            cursor: pointer;
        }
        .customer-section {
            margin-top: 20px;
            padding: 20px;
            border: 1px solid #ddd;
            border-radius: 5px;
        }
        label { font-weight: bold; display: inline-block; width: 80px; }
        input[type="text"], input[type="number"] {
            padding: 8px;
            border: 1px solid #ccc;
            border-radius: 5px;
        }
        .add-link {
            display: inline-block;
            margin-top: 10px;
            padding: 8px 20px;
            background-color: #4CAF50;
            color: white;
            text-decoration: none;
            border-radius: 5px;
        }
        .logout-link {
            display: block;
            text-align: center;
            margin-top: 20px;
            color: #f44336;
            font-weight: bold;
        }
    </style>
</head>
<body>
    <div class="container">
        <h1>Welcome <%%= username %></h1>
        <div class="menu-section">
            <h3>Our Menu Card (Click for a Surprise!)</h3>
            <a href="/surprise">
                <img src="/images/menu.jpg" alt="Menu Card">
            </a>
        </div>
        <div class="customer-section">
            <h3>Add Customer Information</h3>
            <div class="form-group">
                <label for="age">Age:</label>
                <input type="number" id="age" name="age" placeholder="Enter age">
            </div>
            <div class="form-group">
                <label for="mobile">Mobile:</label>
                <input type="text" id="mobile" name="mobile" placeholder="Enter mobile">
            </div>
            <a href="#" class="add-link" onclick="addCustomer()">Add</a>
        </div>
        <a href="/logout" class="logout-link">Logout</a>
    </div>
    <script>
        function addCustomer() {
            var age = document.getElementById('age').value;
            var mobile = document.getElementById('mobile').value;
            if (age && mobile) {
                window.location.href = '/customer?age=' + age + '&mobile=' + mobile;
            } else {
                alert('Please enter both age and mobile.');
            }
        }
    </script>
</body>
</html>""", "views/home.ejs")

# surprise.ejs
add_code_block(doc, """<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Restaurant CRM - Surprise Offer</title>
    <style>
        body {
            font-family: Arial, sans-serif;
            background-color: #FFF3E0;
            display: flex;
            justify-content: center;
            align-items: center;
            height: 100vh;
            margin: 0;
        }
        .surprise-container {
            background: white;
            padding: 40px;
            border-radius: 10px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.1);
            text-align: center;
        }
        h2 { color: #FF5722; }
        h3 { color: #E91E63; font-size: 24px; }
        .logout-link {
            display: block;
            margin-top: 20px;
            color: #f44336;
            font-weight: bold;
        }
    </style>
</head>
<body>
    <div class="surprise-container">
        <h2>Congratulations!</h2>
        <h3>Surprise offer of Rs.250 in the next order</h3>
        <a href="/logout" class="logout-link">Logout</a>
    </div>
</body>
</html>""", "views/surprise.ejs")

# customer.ejs
add_code_block(doc, """<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Restaurant CRM - Customer Information</title>
    <style>
        body {
            font-family: Arial, sans-serif;
            background-color: #f4f4f4;
            padding: 20px;
        }
        .container {
            max-width: 600px;
            margin: 0 auto;
            background: white;
            padding: 30px;
            border-radius: 10px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.1);
        }
        h2 { text-align: center; color: #333; }
        table {
            width: 100%;
            border-collapse: collapse;
            margin-top: 20px;
        }
        th, td {
            border: 1px solid #ddd;
            padding: 12px;
            text-align: center;
        }
        th { background-color: #FF5722; color: white; }
        tr:nth-child(even) { background-color: #f9f9f9; }
        .logout-link {
            display: block;
            text-align: center;
            margin-top: 20px;
            color: #f44336;
            font-weight: bold;
        }
    </style>
</head>
<body>
    <div class="container">
        <h2>Customer Information</h2>
        <%% if (customers.length > 0) { %>
            <table>
                <thead>
                    <tr>
                        <th>S.No</th>
                        <th>Age</th>
                        <th>Mobile</th>
                    </tr>
                </thead>
                <tbody>
                    <%% customers.forEach(function(customer, index) { %>
                        <tr>
                            <td><%%= index + 1 %></td>
                            <td><%%= customer.age %></td>
                            <td><%%= customer.mobile %></td>
                        </tr>
                    <%% }); %>
                </tbody>
            </table>
        <%% } else { %>
            <p style="text-align:center;color:#999;">No customer data available.</p>
        <%% } %>
        <a href="/logout" class="logout-link">Logout</a>
    </div>
</body>
</html>""", "views/customer.ejs")

# --- FILES ---
add_heading_styled(doc, "6. Files", level=2)
files_table2 = doc.add_table(rows=7, cols=3)
files_table2.style = 'Table Grid'
fh = files_table2.rows[0].cells
fh[0].text = "S.No"
fh[1].text = "File Name"
fh[2].text = "Description"
set_cell_shading(fh[0], '003366')
set_cell_shading(fh[1], '003366')
set_cell_shading(fh[2], '003366')
for cell in fh:
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True

fd2 = [
    ("1", "package.json", "Project configuration with dependencies (express, express-session, ejs)"),
    ("2", "app.js", "Main server file with all routes, session config, and middleware"),
    ("3", "views/login.ejs", "Login page template with username and password form"),
    ("4", "views/home.ejs", "Home page with welcome message, menu image, customer form, logout"),
    ("5", "views/surprise.ejs", "Surprise offer page with Rs.250 discount message"),
    ("6", "views/customer.ejs", "Customer page displaying age and mobile in a table"),
]
for i, (sno, fname, desc) in enumerate(fd2):
    row = files_table2.rows[i + 1].cells
    row[0].text = sno
    row[1].text = fname
    row[2].text = desc

doc.add_paragraph("")

# Folder Structure
p = doc.add_paragraph()
run = p.add_run("Folder Structure:")
run.bold = True
add_code_block(doc, """crm-restaurant/
├── app.js
├── package.json
├── public/
│   └── images/
│       └── menu.jpg
└── views/
    ├── login.ejs
    ├── home.ejs
    ├── surprise.ejs
    └── customer.ejs""")

# --- HOW TO RUN ---
add_heading_styled(doc, "7. How to Run", level=2)
add_numbered_list(doc, [
    "Ensure Node.js (v14+) and npm are installed.",
    "Open terminal and navigate to the project folder:\n   cd crm-restaurant",
    "Install dependencies:\n   npm install",
    "Place a menu card image as public/images/menu.jpg (or use any food image).",
    "Start the server:\n   node app.js",
    "Open browser and go to:\n   http://localhost:3000",
    "Login with credentials:\n   Username: admin\n   Password: password123",
    "On the home page, click the menu image to see the surprise offer.",
    "Enter customer age and mobile, click 'Add' to see the customer table.",
    "Click 'Logout' on any page to destroy session and return to login."
])

add_heading_styled(doc, "8. Output", level=2)
add_image_with_caption(doc, r"c:\RISK_ASSESSMENT\genai\fullstack\screenshots\crm_login.png", "Login Page")
add_image_with_caption(doc, r"c:\RISK_ASSESSMENT\genai\fullstack\screenshots\crm_home.png", "Home Page")
add_image_with_caption(doc, r"c:\RISK_ASSESSMENT\genai\fullstack\screenshots\crm_surprise.png", "Surprise Offer Page")
add_image_with_caption(doc, r"c:\RISK_ASSESSMENT\genai\fullstack\screenshots\crm_customer.png", "Customer Page with Table")

add_page_break(doc)

# ================================================================
# SET B - QUESTION 1: Kids Library Angular App
# ================================================================

add_heading_styled(doc, "SET B - QUESTION 1", level=1)
add_heading_styled(doc, "Kids Library Gift Coupon System using Angular", level=2)

# --- AIM ---
add_heading_styled(doc, "1. Aim", level=2)
doc.add_paragraph(
    "To create an Angular application with a service that serves children data for a kid's library. "
    "The application filters children based on their birth month and displays them in two child components: "
    "one for children born in months with 30 days (including February) showing 'Get one book free', "
    "and another for children born in months with 31 days showing 'Get coloring kits free'. "
    "The parent component displays both child components in sequence."
)

# --- ALGORITHM ---
add_heading_styled(doc, "2. Algorithm", level=2)
add_numbered_list(doc, [
    "START",
    "Create a ChildrenService with an array of children data (name, monthBorn, age).",
    "Define two arrays: thirtyDayMonths (Feb, Apr, Jun, Sep, Nov) and thirtyOneDayMonths (Jan, Mar, May, Jul, Aug, Oct, Dec).",
    "Implement getChildrenBornIn30DayMonths() method to filter children born in 30-day months.",
    "Implement getChildrenBornIn31DayMonths() method to filter children born in 31-day months.",
    "Create ThirtyComponent: inject ChildrenService, call getChildrenBornIn30DayMonths(), display names in unordered list with 'Get one book free' message.",
    "Create ThirtyOneComponent: inject ChildrenService, call getChildrenBornIn31DayMonths(), display names in unordered list with 'Get coloring kits free' message.",
    "Create parent AppComponent that includes <app-thirty> followed by <app-thirtyone>.",
    "Register all components and the service in AppModule.",
    "STOP"
])

# --- WORKFLOW ---
add_heading_styled(doc, "3. Workflow", level=2)
add_numbered_list(doc, [
    "User opens the application in the browser.",
    "Angular bootstraps AppComponent (parent).",
    "Parent component renders ThirtyComponent first.",
    "ThirtyComponent calls ChildrenService.getChildrenBornIn30DayMonths().",
    "Children born in Feb, Apr, Jun, Sep, Nov are displayed in an unordered list with 'Get one book free'.",
    "Parent component then renders ThirtyOneComponent.",
    "ThirtyOneComponent calls ChildrenService.getChildrenBornIn31DayMonths().",
    "Children born in Jan, Mar, May, Jul, Aug, Oct, Dec are displayed in an unordered list with 'Get coloring kits free'."
])

# --- PROCEDURE ---
add_heading_styled(doc, "4. Procedure", level=2)
add_numbered_list(doc, [
    "Create a new Angular project: ng new kids-library",
    "Navigate to the project: cd kids-library",
    "Generate the service: ng generate service children",
    "Generate the thirty component: ng generate component thirty",
    "Generate the thirtyone component: ng generate component thirtyone",
    "Implement the ChildrenService with children data and filter methods.",
    "Implement ThirtyComponent to consume the service and display 30-day month children.",
    "Implement ThirtyOneComponent to consume the service and display 31-day month children.",
    "Update the parent AppComponent template to include both child components.",
    "Register everything in AppModule.",
    "Run the application: ng serve",
    "Open browser at http://localhost:4200 to verify output."
])

# --- CODE ---
add_heading_styled(doc, "5. Code", level=2)

# children.service.ts
add_code_block(doc, """import { Injectable } from '@angular/core';

export interface Child {
  name: string;
  monthBorn: string;
  age: number;
}

@Injectable({
  providedIn: 'root'
})
export class ChildrenService {
  private children: Child[] = [
    { name: 'Alice', monthBorn: 'January', age: 7 },
    { name: 'Bob', monthBorn: 'February', age: 5 },
    { name: 'Charlie', monthBorn: 'March', age: 8 },
    { name: 'Diana', monthBorn: 'April', age: 6 },
    { name: 'Eve', monthBorn: 'May', age: 9 },
    { name: 'Frank', monthBorn: 'June', age: 4 },
    { name: 'Grace', monthBorn: 'July', age: 7 },
    { name: 'Henry', monthBorn: 'August', age: 10 },
    { name: 'Ivy', monthBorn: 'September', age: 6 },
    { name: 'Jack', monthBorn: 'October', age: 8 },
    { name: 'Kate', monthBorn: 'November', age: 5 },
    { name: 'Leo', monthBorn: 'December', age: 7 }
  ];

  // Months with 30 days (including February as per requirement)
  private thirtyDayMonths: string[] = [
    'February', 'April', 'June', 'September', 'November'
  ];

  // Months with 31 days
  private thirtyOneDayMonths: string[] = [
    'January', 'March', 'May', 'July', 'August', 'October', 'December'
  ];

  getAllChildren(): Child[] {
    return this.children;
  }

  getChildrenBornIn30DayMonths(): Child[] {
    return this.children.filter(child =>
      this.thirtyDayMonths.includes(child.monthBorn)
    );
  }

  getChildrenBornIn31DayMonths(): Child[] {
    return this.children.filter(child =>
      this.thirtyOneDayMonths.includes(child.monthBorn)
    );
  }
}""", "src/app/children.service.ts")

# thirty.component.ts
add_code_block(doc, """import { Component, OnInit } from '@angular/core';
import { ChildrenService, Child } from '../children.service';

@Component({
  selector: 'app-thirty',
  templateUrl: './thirty.component.html',
  styleUrls: ['./thirty.component.css']
})
export class ThirtyComponent implements OnInit {
  children: Child[] = [];

  constructor(private childrenService: ChildrenService) {}

  ngOnInit(): void {
    this.children = this.childrenService.getChildrenBornIn30DayMonths();
  }
}""", "src/app/thirty/thirty.component.ts")

# thirty.component.html
add_code_block(doc, """<div class="card">
  <h2>Children Born in Months with 30 Days (Including February)</h2>
  <p class="offer">Get one book free</p>
  <ul>
    <li *ngFor="let child of children">
      {{ child.name }} ({{ child.monthBorn }}, Age: {{ child.age }})
    </li>
  </ul>
</div>""", "src/app/thirty/thirty.component.html")

# thirty.component.css
add_code_block(doc, """.card {
  border: 2px solid #4CAF50;
  border-radius: 10px;
  padding: 20px;
  margin: 10px;
  background-color: #f0fff0;
}
h2 { color: #2E7D32; }
.offer {
  font-size: 18px;
  font-weight: bold;
  color: #FF5722;
  background-color: #FFF3E0;
  padding: 10px;
  border-radius: 5px;
  text-align: center;
}
ul { list-style-type: disc; padding-left: 20px; }
li { padding: 5px 0; font-size: 16px; }""", "src/app/thirty/thirty.component.css")

# thirtyone.component.ts
add_code_block(doc, """import { Component, OnInit } from '@angular/core';
import { ChildrenService, Child } from '../children.service';

@Component({
  selector: 'app-thirtyone',
  templateUrl: './thirtyone.component.html',
  styleUrls: ['./thirtyone.component.css']
})
export class ThirtyOneComponent implements OnInit {
  children: Child[] = [];

  constructor(private childrenService: ChildrenService) {}

  ngOnInit(): void {
    this.children = this.childrenService.getChildrenBornIn31DayMonths();
  }
}""", "src/app/thirtyone/thirtyone.component.ts")

# thirtyone.component.html
add_code_block(doc, """<div class="card">
  <h2>Children Born in Months with 31 Days</h2>
  <p class="offer">Get coloring kits free</p>
  <ul>
    <li *ngFor="let child of children">
      {{ child.name }} ({{ child.monthBorn }}, Age: {{ child.age }})
    </li>
  </ul>
</div>""", "src/app/thirtyone/thirtyone.component.html")

# thirtyone.component.css
add_code_block(doc, """.card {
  border: 2px solid #1565C0;
  border-radius: 10px;
  padding: 20px;
  margin: 10px;
  background-color: #E3F2FD;
}
h2 { color: #0D47A1; }
.offer {
  font-size: 18px;
  font-weight: bold;
  color: #E91E63;
  background-color: #FCE4EC;
  padding: 10px;
  border-radius: 5px;
  text-align: center;
}
ul { list-style-type: disc; padding-left: 20px; }
li { padding: 5px 0; font-size: 16px; }""", "src/app/thirtyone/thirtyone.component.css")

# app.component.ts
add_code_block(doc, """import { Component } from '@angular/core';

@Component({
  selector: 'app-root',
  templateUrl: './app.component.html',
  styleUrls: ['./app.component.css']
})
export class AppComponent {
  title = "Kid's Library - Gift Coupon System";
}""", "src/app/app.component.ts")

# app.component.html
add_code_block(doc, """<div class="container">
  <h1>{{ title }}</h1>
  <app-thirty></app-thirty>
  <app-thirtyone></app-thirtyone>
</div>""", "src/app/app.component.html")

# app.component.css
add_code_block(doc, """.container {
  max-width: 600px;
  margin: 30px auto;
  padding: 20px;
  font-family: Arial, sans-serif;
}
h1 {
  text-align: center;
  color: #333;
  margin-bottom: 20px;
}""", "src/app/app.component.css")

# app.module.ts
add_code_block(doc, """import { NgModule } from '@angular/core';
import { BrowserModule } from '@angular/platform-browser';
import { AppComponent } from './app.component';
import { ThirtyComponent } from './thirty/thirty.component';
import { ThirtyOneComponent } from './thirtyone/thirtyone.component';
import { ChildrenService } from './children.service';

@NgModule({
  declarations: [
    AppComponent,
    ThirtyComponent,
    ThirtyOneComponent
  ],
  imports: [BrowserModule],
  providers: [ChildrenService],
  bootstrap: [AppComponent]
})
export class AppModule {}""", "src/app/app.module.ts")

# main.ts
add_code_block(doc, """import { platformBrowserDynamic } from '@angular/platform-browser-dynamic';
import { AppModule } from './app/app.module';

platformBrowserDynamic().bootstrapModule(AppModule)
  .catch(err => console.error(err));""", "src/main.ts")

# index.html
add_code_block(doc, """<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <title>Kids Library - Gift Coupon System</title>
  <base href="/">
  <meta name="viewport" content="width=device-width, initial-scale=1">
</head>
<body>
  <app-root></app-root>
</body>
</html>""", "src/index.html")

# --- FILES ---
add_heading_styled(doc, "6. Files", level=2)
files_table3 = doc.add_table(rows=13, cols=3)
files_table3.style = 'Table Grid'
fh3 = files_table3.rows[0].cells
fh3[0].text = "S.No"
fh3[1].text = "File Name"
fh3[2].text = "Description"
set_cell_shading(fh3[0], '003366')
set_cell_shading(fh3[1], '003366')
set_cell_shading(fh3[2], '003366')
for cell in fh3:
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True

fd3 = [
    ("1", "src/app/children.service.ts", "Service with children data and filter methods"),
    ("2", "src/app/thirty/thirty.component.ts", "Component logic for 30-day month children"),
    ("3", "src/app/thirty/thirty.component.html", "Template displaying 30-day month children list"),
    ("4", "src/app/thirty/thirty.component.css", "Styles for thirty component (green theme)"),
    ("5", "src/app/thirtyone/thirtyone.component.ts", "Component logic for 31-day month children"),
    ("6", "src/app/thirtyone/thirtyone.component.html", "Template displaying 31-day month children list"),
    ("7", "src/app/thirtyone/thirtyone.component.css", "Styles for thirtyone component (blue theme)"),
    ("8", "src/app/app.component.ts", "Parent component definition"),
    ("9", "src/app/app.component.html", "Parent template with both child components"),
    ("10", "src/app/app.component.css", "Parent component styles"),
    ("11", "src/app/app.module.ts", "Root module registering all components and service"),
    ("12", "src/main.ts", "Bootstrap file for the Angular application"),
]
for i, (sno, fname, desc) in enumerate(fd3):
    row = files_table3.rows[i + 1].cells
    row[0].text = sno
    row[1].text = fname
    row[2].text = desc

# --- HOW TO RUN ---
add_heading_styled(doc, "7. How to Run", level=2)
add_numbered_list(doc, [
    "Ensure Node.js (v14+) and npm are installed.",
    "Install Angular CLI: npm install -g @angular/cli",
    "Create a new project: ng new kids-library",
    "Navigate to the project: cd kids-library",
    "Generate the service: ng generate service children",
    "Generate components:\n   ng generate component thirty\n   ng generate component thirtyone",
    "Replace the generated files with the code provided above.",
    "Run the application: ng serve",
    "Open browser at http://localhost:4200",
    "Verify: Children born in 30-day months appear first with 'Get one book free'.",
    "Verify: Children born in 31-day months appear below with 'Get coloring kits free'."
])

doc.add_paragraph("")
p = doc.add_paragraph()
run = p.add_run("Expected Output:")
run.bold = True
run.font.size = Pt(11)
doc.add_paragraph("30-Day Months Section: Bob (February), Diana (April), Frank (June), Ivy (September), Kate (November)")
doc.add_paragraph("31-Day Months Section: Alice (January), Charlie (March), Eve (May), Grace (July), Henry (August), Jack (October), Leo (December)")

add_heading_styled(doc, "8. Output", level=2)
add_image_with_caption(doc, r"c:\RISK_ASSESSMENT\genai\fullstack\screenshots\setB_q1_output.png", "Kids Library Gift Coupon Output")

add_page_break(doc)

# ================================================================
# SET B - QUESTION 2: CRM Restaurant (same as SET A Q2)
# ================================================================

add_heading_styled(doc, "SET B - QUESTION 2", level=1)
add_heading_styled(doc, "CRM Application for Restaurant using Node.js and Express", level=2)

doc.add_paragraph("")
p = doc.add_paragraph()
run = p.add_run("Note: ")
run.bold = True
run.font.color.rgb = RGBColor(200, 0, 0)
run2 = p.add_run(
    "This question is identical to SET A - Question 2. The complete solution including "
    "Aim, Algorithm, Workflow, Procedure, Code, Files, and How to Run is provided in "
    "SET A - Question 2 (pages above). Please refer to SET A - Question 2 for the full solution."
)

doc.add_paragraph("")
doc.add_paragraph("For quick reference, the key details are repeated below:")

# --- AIM ---
add_heading_styled(doc, "1. Aim", level=2)
doc.add_paragraph(
    "To design and develop a Customer Relationship Management (CRM) application for a restaurant "
    "using Node.js and Express framework with session management, login authentication, menu card "
    "display, surprise offers, customer information management, and logout functionality."
)

# --- ALGORITHM ---
add_heading_styled(doc, "2. Algorithm", level=2)
add_numbered_list(doc, [
    "START",
    "Initialize Express app with session middleware and EJS view engine.",
    "Define hardcoded credentials (username: admin, password: password123).",
    "GET /login → Render login page with username and password fields.",
    "POST /home → Validate credentials, create session, render home page with 'Welcome <<username>>'.",
    "Home page has: menu image (hyperlink to /surprise), customer form (age, mobile, 'add' link).",
    "GET /surprise → Display 'surprise offer of Rs.250 in the next order' in <h3> tag.",
    "GET /customer → Add customer (age, mobile) to array, display all in table.",
    "GET /logout → Destroy session, redirect to /login.",
    "All pages have logout hyperlink.",
    "STOP"
])

# --- WORKFLOW ---
add_heading_styled(doc, "3. Workflow", level=2)
add_numbered_list(doc, [
    "User accesses http://localhost:3000 → redirected to /login.",
    "Enter credentials → POST /home → session created → Welcome page.",
    "Click menu image → /surprise → Rs.250 offer displayed.",
    "Enter customer info → click 'Add' → /customer → table displayed.",
    "Click 'Logout' → session destroyed → redirect to /login."
])

# --- PROCEDURE ---
add_heading_styled(doc, "4. Procedure", level=2)
doc.add_paragraph("Same as SET A - Question 2. Refer to the detailed procedure above.")

# --- CODE ---
add_heading_styled(doc, "5. Code", level=2)
doc.add_paragraph("The complete code is identical to SET A - Question 2. The six files are:")
add_bullet_list(doc, [
    "package.json - Project configuration",
    "app.js - Main server with all routes and session management",
    "views/login.ejs - Login page template",
    "views/home.ejs - Home page with welcome, menu image, customer form",
    "views/surprise.ejs - Surprise offer page",
    "views/customer.ejs - Customer information table page"
])
doc.add_paragraph("Refer to SET A - Question 2 for the complete code listings.")

# --- FILES ---
add_heading_styled(doc, "6. Files", level=2)
doc.add_paragraph("Same file structure as SET A - Question 2:")
add_code_block(doc, """crm-restaurant/
├── app.js
├── package.json
├── public/
│   └── images/
│       └── menu.jpg
└── views/
    ├── login.ejs
    ├── home.ejs
    ├── surprise.ejs
    └── customer.ejs""")

# --- HOW TO RUN ---
add_heading_styled(doc, "7. How to Run", level=2)
add_numbered_list(doc, [
    "cd crm-restaurant",
    "npm install",
    "node app.js",
    "Open http://localhost:3000",
    "Login: admin / password123",
    "Test all features: menu image, surprise offer, customer add, logout."
])

# ================================================================
# Save the document
# ================================================================

output_path = r"c:\RISK_ASSESSMENT\genai\fullstack\Fullstack_Lab_Exam_Document.docx"
doc.save(output_path)
print(f"Document saved successfully at: {output_path}")
print(f"Total sections: 4 questions covered (SET A Q1, SET A Q2, SET B Q1, SET B Q2)")
